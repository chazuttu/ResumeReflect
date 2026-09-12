import io
import os
import json
import re

import groq
from fastapi import FastAPI, HTTPException
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import StreamingResponse
from pydantic import BaseModel

from docx import Document
from docx.shared import Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

# ─────────────────────────────────────────────
# CONFIG
# ─────────────────────────────────────────────
GROQ_API_KEY = os.environ.get("GROQ_API_KEY", "")

app = FastAPI(title="ResumeReflect API")

app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],  # tighten this to your app's origin before going live
    allow_methods=["*"],
    allow_headers=["*"],
)

# ─────────────────────────────────────────────
# REQUEST / RESPONSE MODELS
# ─────────────────────────────────────────────
class AnalyzeRequest(BaseModel):
    resume_text: str
    jd_text: str
    market_mode: str = "International (Workday / Greenhouse / Lever)"


class ResumeDocxRequest(BaseModel):
    data: dict
    watermark: bool = False


# ─────────────────────────────────────────────
# CORE LOGIC (ported from app.py, unchanged)
# ─────────────────────────────────────────────
def call_groq(resume_text, jd_text, market_mode):
    if not GROQ_API_KEY:
        raise HTTPException(status_code=500, detail="GROQ_API_KEY is not set on the server")
    client = groq.Groq(api_key=GROQ_API_KEY)

    if "India" in market_mode:
        market_instructions = """
INDIAN JOB MARKET RULES (very important):
- Optimise keywords specifically for Naukri.com and LinkedIn India ATS ranking
- Use Indian resume conventions: include notice period if mentioned, CTC in LPA format, percentage-based education scores
- Add Indian recruiter search terms naturally: "immediate joiner", "open to relocation", relevant Indian tech stack terms
- Keep declaration section if present in original resume
- Bullet points should include measurable Indian industry-standard metrics
- Summary should mention notice period and location preference if available
- Skills should match exactly what Indian recruiters search for on Naukri
"""
    else:
        market_instructions = """
INTERNATIONAL JOB MARKET RULES (very important):
- Optimise keywords for global ATS systems: Workday, Greenhouse, Lever, Indeed
- Use international resume conventions: no photo, no DOB, no declaration, clean 1-page preferred
- Salary references in annual USD/GBP format if mentioned
- Use action verbs and quantified achievements suited for western hiring managers
- Skills and tools should match global industry-standard terminology
- Summary should be punchy, achievement-focused, and ATS-friendly for international roles
"""

    prompt = f"""
You are an expert ATS resume specialist and career coach.
Analyze the resume against the job description carefully.
Return ONLY a JSON object. No text before or after. No markdown. Just pure JSON.

{market_instructions}

{{
  "candidate_name": "full name from resume",
  "email": "email from resume or empty string",
  "phone": "phone from resume or empty string",
  "location": "city from resume or empty string",
  "linkedin": "linkedin url from resume or empty string",
  "match_score": 75,
  "ats_keywords_found": 12,
  "ats_keywords_missing": 5,
  "strong_points": ["point 1","point 2","point 3","point 4","point 5"],
  "missing_skills": ["skill 1","skill 2","skill 3","skill 4"],
  "improvement_tips": ["tip 1","tip 2","tip 3","tip 4"],
  "summary": "2 to 3 sentence professional summary tailored to the job description",
  "work_experience": [{{
    "title": "job title exactly as in resume",
    "company": "company name exactly as in resume",
    "dates": "dates exactly as in resume",
    "location": "location exactly as in resume",
    "bullets": ["bullet rewritten with JD keywords","bullet","bullet"]
  }}],
  "projects": [{{
    "name": "project name",
    "bullets": ["project description bullet"]
  }}],
  "education": [{{
    "degree": "degree name exactly as in resume",
    "institution": "institution name exactly as in resume",
    "year": "year exactly as in resume",
    "cgpa": "cgpa or empty string"
  }}],
  "skills_technical": ["skill1","skill2","skill3"],
  "skills_tools": ["tool1","tool2","tool3"],
  "achievements": ["achievement 1","achievement 2","achievement 3"],
  "certifications": ["certification 1","certification 2"],
  "score_explanation": "2-3 sentences explaining exactly why the ATS score improved — mention specific keywords added, sections strengthened, and what made the biggest difference",
  "job_title_suggestions": ["Job Title 1","Job Title 2","Job Title 3","Job Title 4","Job Title 5"]
}}

STRICT RULES:
- Never fabricate any experience skills or education
- Only use information already present in the resume
- Rewrite bullet points using keywords from the job description
- Keep all dates company names and institutions exactly as original
- match_score must be a number not a string
- Return ONLY pure JSON nothing else

RESUME:
{resume_text}

JOB DESCRIPTION:
{jd_text}
"""
    response = client.chat.completions.create(
        model="llama-3.3-70b-versatile",
        messages=[{"role": "user", "content": prompt}],
        temperature=0.5,
        max_tokens=3500,
    )
    return response.choices[0].message.content


def parse_json(text):
    text = text.strip()
    if "```" in text:
        lines = [l for l in text.split("\n") if not l.strip().startswith("```")]
        text = "\n".join(lines)
    start = text.find("{")
    end = text.rfind("}") + 1
    if start != -1 and end > start:
        text = text[start:end]
    return json.loads(text)


# ─────────────────────────────────────────────
# DOCX HELPERS (ported from app.py, unchanged)
# ─────────────────────────────────────────────
def set_cell_bg(cell, hex_color):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), hex_color)
    tcPr.append(shd)


def set_cell_borders(cell, **kwargs):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcBorders = OxmlElement('w:tcBorders')
    for side in ['top', 'left', 'bottom', 'right']:
        tag = OxmlElement(f'w:{side}')
        tag.set(qn('w:val'), kwargs.get(side, 'nil'))
        tag.set(qn('w:sz'), '0')
        tag.set(qn('w:space'), '0')
        tag.set(qn('w:color'), 'auto')
        tcBorders.append(tag)
    tcPr.append(tcBorders)


def section_heading(doc, text, color="2E75B6"):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(4)
    pPr = p._p.get_or_add_pPr()
    pBdr = OxmlElement('w:pBdr')
    bottom = OxmlElement('w:bottom')
    bottom.set(qn('w:val'), 'single')
    bottom.set(qn('w:sz'), '6')
    bottom.set(qn('w:space'), '1')
    bottom.set(qn('w:color'), color)
    pBdr.append(bottom)
    pPr.append(pBdr)
    run = p.add_run(text.upper())
    run.bold = True
    run.font.size = Pt(10)
    run.font.color.rgb = RGBColor(*bytes.fromhex(color))
    return p


def bullet_para(doc, text, size=9.5):
    p = doc.add_paragraph(style='List Bullet')
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(2)
    run = p.add_run(text.lstrip('•-– '))
    run.font.size = Pt(size)
    return p


def build_resume(data, watermark=False):
    doc = Document()

    for section in doc.sections:
        section.top_margin = Cm(1.5)
        section.bottom_margin = Cm(1.5)
        section.left_margin = Cm(1.8)
        section.right_margin = Cm(1.8)

    doc.styles['Normal'].font.name = 'Calibri'
    doc.styles['Normal'].font.size = Pt(10)

    if watermark:
        hdr = doc.sections[0].header
        hp = hdr.paragraphs[0]
        hp.clear()
        run = hp.add_run("FREE VERSION — UPGRADE FOR CLEAN COPY")
        run.font.size = Pt(7)
        run.font.color.rgb = RGBColor(0x9B, 0x59, 0xB6)
        hp.alignment = WD_ALIGN_PARAGRAPH.CENTER

    name_tbl = doc.add_table(rows=1, cols=1)
    name_tbl.style = 'Table Grid'
    cell = name_tbl.cell(0, 0)
    set_cell_bg(cell, '1A1A2E')
    set_cell_borders(cell, top='nil', bottom='nil', left='nil', right='nil')
    cell.paragraphs[0].clear()

    np = cell.paragraphs[0]
    nr = np.add_run((data.get('candidate_name') or 'Candidate').upper())
    nr.bold = True
    nr.font.size = Pt(18)
    nr.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
    np.paragraph_format.space_before = Pt(8)
    np.paragraph_format.space_after = Pt(2)

    contact_parts = [x for x in [
        data.get('email'), data.get('phone'),
        data.get('location'), data.get('linkedin')
    ] if x]
    cp = cell.add_paragraph("   |   ".join(contact_parts))
    cp.runs[0].font.size = Pt(8)
    cp.runs[0].font.color.rgb = RGBColor(0xA0, 0xC4, 0xFF)
    cp.paragraph_format.space_before = Pt(0)
    cp.paragraph_format.space_after = Pt(8)

    doc.add_paragraph()

    if data.get('summary'):
        section_heading(doc, 'Professional Summary')
        p = doc.add_paragraph(data['summary'])
        p.runs[0].italic = True
        p.runs[0].font.size = Pt(9.5)
        p.paragraph_format.space_after = Pt(4)

    jobs = data.get('work_experience', [])
    if jobs:
        section_heading(doc, 'Work Experience')
        for job in jobs:
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(4)
            p.paragraph_format.space_after = Pt(1)
            r1 = p.add_run(job.get('title', ''))
            r1.bold = True
            r1.font.size = Pt(10)
            r2 = p.add_run(f"  ·  {job.get('company', '')}")
            r2.bold = True
            r2.font.size = Pt(10)
            r2.font.color.rgb = RGBColor(0x2E, 0x75, 0xB6)
            r3 = p.add_run(f"  ·  {job.get('dates', '')}  ·  {job.get('location', '')}")
            r3.italic = True
            r3.font.size = Pt(8.5)
            r3.font.color.rgb = RGBColor(0x66, 0x66, 0x66)
            for b in job.get('bullets', [])[:3]:
                bullet_para(doc, b)

    projects = data.get('projects', [])
    if projects:
        section_heading(doc, 'Projects')
        for proj in projects[:2]:
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(4)
            p.paragraph_format.space_after = Pt(1)
            r = p.add_run(proj.get('name', ''))
            r.bold = True
            r.font.size = Pt(10)
            r.font.color.rgb = RGBColor(0x2E, 0x75, 0xB6)
            for b in proj.get('bullets', [])[:2]:
                bullet_para(doc, b)

    tech = data.get('skills_technical', [])
    tools = data.get('skills_tools', [])
    if tech or tools:
        section_heading(doc, 'Skills')
        skills_tbl = doc.add_table(rows=1, cols=2)
        skills_tbl.style = 'Table Grid'

        lc = skills_tbl.cell(0, 0)
        rc = skills_tbl.cell(0, 1)
        set_cell_bg(lc, 'F5F7FA')
        set_cell_bg(rc, 'F5F7FA')
        set_cell_borders(lc, top='nil', bottom='nil', left='nil', right='nil')
        set_cell_borders(rc, top='nil', bottom='nil', left='nil', right='nil')

        lc.paragraphs[0].clear()
        lp = lc.paragraphs[0]
        lr = lp.add_run('Technical Skills')
        lr.bold = True
        lr.font.size = Pt(9)
        lr.font.color.rgb = RGBColor(0x2E, 0x75, 0xB6)
        lp2 = lc.add_paragraph("  •  ".join(tech))
        lp2.runs[0].font.size = Pt(9)

        rc.paragraphs[0].clear()
        rp = rc.paragraphs[0]
        rr = rp.add_run('Tools & Technologies')
        rr.bold = True
        rr.font.size = Pt(9)
        rr.font.color.rgb = RGBColor(0x2E, 0x75, 0xB6)
        rp2 = rc.add_paragraph("  •  ".join(tools))
        rp2.runs[0].font.size = Pt(9)

    edu = data.get('education', [])
    if edu:
        section_heading(doc, 'Education')
        for e in edu:
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(3)
            p.paragraph_format.space_after = Pt(2)
            r1 = p.add_run(e.get('degree', ''))
            r1.bold = True
            r1.font.size = Pt(10)
            extra = f"  ·  {e.get('institution', '')}  ·  {e.get('year', '')}"
            if e.get('cgpa'):
                extra += f"  ·  CGPA: {e['cgpa']}"
            r2 = p.add_run(extra)
            r2.font.size = Pt(9)
            r2.font.color.rgb = RGBColor(0x66, 0x66, 0x66)

    ach = data.get('achievements', [])
    if ach:
        section_heading(doc, 'Achievements')
        for a in ach[:3]:
            bullet_para(doc, a)

    certs = data.get('certifications', [])
    if certs:
        section_heading(doc, 'Certifications')
        for c in certs[:3]:
            bullet_para(doc, c)

    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(10)
    pPr = p._p.get_or_add_pPr()
    pBdr = OxmlElement('w:pBdr')
    top = OxmlElement('w:top')
    top.set(qn('w:val'), 'single')
    top.set(qn('w:sz'), '2')
    top.set(qn('w:space'), '1')
    top.set(qn('w:color'), 'D0D8E8')
    pBdr.append(top)
    pPr.append(pBdr)
    run = p.add_run("Tailored with ResumeReflect  •  AI Resume Tool")
    run.italic = True
    run.font.size = Pt(7.5)
    run.font.color.rgb = RGBColor(0xAA, 0xAA, 0xAA)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER

    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)
    return buf.read()


# ─────────────────────────────────────────────
# ROUTES
# ─────────────────────────────────────────────
@app.get("/health")
def health():
    return {"status": "ok"}


@app.post("/api/analyze")
def analyze(req: AnalyzeRequest):
    if len(req.resume_text.strip()) < 30 or len(req.jd_text.strip()) < 30:
        raise HTTPException(status_code=400, detail="Resume and job description are too short")
    raw = call_groq(req.resume_text, req.jd_text, req.market_mode)
    try:
        data = parse_json(raw)
    except Exception:
        raise HTTPException(status_code=502, detail="AI response could not be parsed, try again")
    return data


@app.post("/api/resume-docx")
def resume_docx(req: ResumeDocxRequest):
    file_bytes = build_resume(req.data, watermark=req.watermark)
    return StreamingResponse(
        io.BytesIO(file_bytes),
        media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        headers={"Content-Disposition": "attachment; filename=resumereflect_tailored_resume.docx"},
    )
